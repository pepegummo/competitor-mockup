"""ส่งออกรายงาน (markdown ที่ LLM เขียน + ลิงก์อ้างอิง) เป็นไฟล์ Word ที่ลิงก์กดได้"""
import io
import re

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Tahoma"  # มีทุกเครื่อง Windows และแสดงภาษาไทยได้
INLINE = re.compile(r"\[\[(\d+)\]\]\(([^)\s]+)\)|\[([^\]]+)\]\(([^)\s]+)\)|\*\*([^*]+)\*\*")
LINK_COLOR = RGBColor(0x25, 0x57, 0xA7)


def _font(run_or_style, size=None):
    run_or_style.font.name = FONT
    rpr = run_or_style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):  # cs = complex script (ภาษาไทย)
        fonts.set(qn(attr), FONT)
    if size:
        run_or_style.font.size = Pt(size)


def _link(paragraph, text, url):
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = paragraph.add_run(text)
    run.font.color.rgb = LINK_COLOR
    run.font.underline = True
    _font(run)
    link.append(run._r)
    paragraph._p.append(link)


def _inline(paragraph, text):
    """**ตัวหนา**, [ข้อความ](url), [[3]](url) → run ธรรมดา/ตัวหนา/ลิงก์"""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _font(paragraph.add_run(text[pos:m.start()]))
        if m.group(1):
            _link(paragraph, f"[{m.group(1)}]", m.group(2))
        elif m.group(3):
            _link(paragraph, m.group(3), m.group(4))
        else:
            run = paragraph.add_run(m.group(5))
            run.bold = True
            _font(run)
        pos = m.end()
    if pos < len(text):
        _font(paragraph.add_run(text[pos:]))


def report_docx(markdown, title, subtitle=""):
    """คืน bytes ของไฟล์ .docx"""
    doc = Document()
    _font(doc.styles["Normal"], 11)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        _font(doc.styles[name])
    doc.add_heading(title, level=0)
    if subtitle:
        _inline(doc.add_paragraph(), subtitle)
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            continue
        head = re.match(r"^(#{1,4})\s+(.*)", line)
        bullet = re.match(r"^\s*[-*•]\s+(.*)", line)
        number = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if head:
            p = doc.add_heading(level=min(len(head.group(1)), 3))
            _inline(p, head.group(2).replace("**", ""))
        elif bullet:
            _inline(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        elif number:
            _inline(doc.add_paragraph(style="List Number"), number.group(1))
        else:
            _inline(doc.add_paragraph(), line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
